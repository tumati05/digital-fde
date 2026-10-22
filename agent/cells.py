"""Digital FDE agent: five step cells. Each turn = 1 model call; manual = +1 call."""

N1 = r'''
import os, re, json, tempfile, html, httpx
from urllib.parse import urlparse
import pandas as pd
from IPython.display import display, HTML

NL = chr(10)
AREAS = [('business_goal','Business goal'),('current_process','Current process'),('systems_involved','Systems involved'),('volume_per_day','Volume'),('exception_handling','Exceptions'),('security_compliance','Security and compliance'),('success_metrics','Success measures'),('users_and_roles','Users and roles'),('data_sources_and_formats','Data and documents'),('integrations_required','Integrations'),('timeline_constraints','Timeline and constraints'),('stakeholders_approvers','Approvers')]
LABELS = dict(AREAS)
CRITICAL = ['business_goal','current_process','systems_involved','volume_per_day','exception_handling','security_compliance','success_metrics','stakeholders_approvers']
AGENT_VERSION = '2026-10-22.4'
QUESTIONS = {
 'business_goal': 'What outcome is the client after, in a sentence, and what makes it urgent now?',
 'current_process': 'Can you walk me through how this work gets done today, from the moment it arrives to the moment it is finished?',
 'systems_involved': 'Which systems or tools does the work touch today?',
 'volume_per_day': 'Roughly how many items a day does the team handle, and when are the busy peaks?',
 'exception_handling': 'What happens today when something does not fit the normal flow, and who steps in?',
 'security_compliance': 'Are there security, privacy or compliance rules to follow, such as where data may live or what needs an audit trail?',
 'success_metrics': 'How will the client judge success: time saved, fewer errors, lower cost, or something else?',
 'users_and_roles': 'Who will work with this day to day, and what are their roles?',
 'data_sources_and_formats': 'What comes in, and in what shape: emails, PDFs, spreadsheets, forms?',
 'integrations_required': 'Which systems would this need to read from or write to, and do they have APIs?',
 'timeline_constraints': 'Is there a deadline, budget limit or other constraint to plan around?',
 'stakeholders_approvers': 'Who signs off on scope, and who approves going live?'}
RATE = {'DISCOVERY_WORKSHOP':16,'PROCESS_MAPPING':24,'INTEGRATION_SIMPLE':24,'INTEGRATION_COMPLEX':60,'DOCUMENT_EXTRACTION_AUTO':40,'WORKFLOW_AUTO_BUILD':48,'DATA_CLEANUP_MIGRATION':40,'POLICY_RULESET':24,'HUMAN_REVIEW_STEP':16,'REPORTING_DASHBOARD':32,'SECURITY_REVIEW':24,'UAT_AND_HYPERCARE':40,'TRAINING_HANDOVER':16}
RATE_SOURCE = 'starter'
RATE_CARD_BAD = False
_rc = str(user_inputs.get('rate_card_json') or '').strip()
if _rc:
    try:
        _o = json.loads(_rc)
        _new = {k: int(v) for k, v in _o.items() if k in RATE and 0 < int(v) <= 1000}
        if _new:
            RATE.update(_new)
            RATE_SOURCE = 'supplied' if len(_new) == len(RATE) else 'partly supplied'
        else:
            RATE_CARD_BAD = True
    except Exception:
        RATE_CARD_BAD = True
MAXV, MAXE = 400, 160
PROV = ('FACT', 'INFERENCE', 'ASSUMPTION')

def n1_state(raw):
    st = {'client_name': '', 'fields': {}, 'wbs': [], 'risks': [], 'turn': 0}
    raw = str(raw or '').strip()
    if not raw:
        return st, False
    if len(raw) > 14000:
        return st, True
    try:
        obj = json.loads(raw)
    except Exception:
        return st, True
    if not isinstance(obj, dict):
        return st, True
    st['client_name'] = str(obj.get('client_name') or '').strip()[:200]
    fields = obj.get('fields') if isinstance(obj.get('fields'), dict) else {}
    for k, _ in AREAS:
        it = fields.get(k)
        if isinstance(it, dict):
            v = str(it.get('v') or '').strip()[:MAXV]
            p = it.get('p')
            if v and p in PROV:
                st['fields'][k] = {'v': v, 'p': p, 'e': str(it.get('e') or '')[:MAXE]}
    wbs = obj.get('wbs') if isinstance(obj.get('wbs'), list) else []
    for w in wbs[:30]:
        if isinstance(w, dict) and w.get('unit') in RATE:
            try:
                q = int(w.get('quantity'))
            except Exception:
                continue
            st['wbs'].append({'unit': w['unit'], 'quantity': min(max(q, 1), 50), 'reason': str(w.get('reason') or '')[:200]})
    risks = obj.get('risks') if isinstance(obj.get('risks'), list) else []
    st['risks'] = [str(r)[:200] for r in risks[:6]]
    try:
        st['turn'] = min(max(int(obj.get('turn') or 0), 0), 500)
    except Exception:
        st['turn'] = 0
    return st, False

def n1_turns(raw):
    raw = str(raw or '').strip()
    if not raw or len(raw) > 20000:
        return []
    try:
        arr = json.loads(raw)
    except Exception:
        return []
    out = []
    if isinstance(arr, list):
        for t in arr[-6:]:
            if isinstance(t, dict) and t.get('role') in ('user', 'assistant'):
                out.append({'role': t['role'], 'text': str(t.get('text') or '')[:600]})
    return out

step_n1_action = str(user_inputs.get('action') or 'chat').strip().lower()
if step_n1_action not in ('chat', 'manual'):
    step_n1_action = 'chat'
step_n1_msg = str(user_inputs.get('user_message') or '').strip()[:2000]
if not step_n1_msg:
    if step_n1_action == 'manual':
        step_n1_msg = 'Please create the manual.'
    else:
        raise ValueError('user_message is empty. Tell me what the client needs.')
step_n1_state, step_n1_reset = n1_state(user_inputs.get('state_json'))
step_n1_turns = n1_turns(user_inputs.get('recent_turns'))
step_n1_notices = []
if RATE_CARD_BAD:
    step_n1_notices.append('The rate card could not be read, so the starter hours were used.')
if step_n1_reset:
    step_n1_notices.append('The saved conversation could not be read, so this chat started fresh.')

step_n1_ref = user_inputs.get('sample_file')
if isinstance(step_n1_ref, list):
    step_n1_ref = step_n1_ref[0] if step_n1_ref else None
step_n1_profile_text = ''

async def n1_localize(ref):
    ref = str(ref)
    if os.path.exists(ref):
        return ref          # a file the platform already stored locally (web app upload)
    u = urlparse(ref)
    host = (u.hostname or '').lower()
    if u.scheme != 'https' or not host or host == 'localhost' or host.endswith('.local') or host.endswith('.internal'):
        raise ValueError('only public https links are allowed')
    import ipaddress
    try:
        ipaddress.ip_address(host)
        raise ValueError('address links are not allowed')
    except ValueError as e:
        if 'not allowed' in str(e):
            raise
    ext = os.path.splitext(u.path)[1].lower()
    async with httpx.AsyncClient(timeout=60, follow_redirects=False) as c:
        async with c.stream('GET', ref) as r:
            if 300 <= r.status_code < 400:
                raise ValueError('redirects are not followed')
            r.raise_for_status()
            data = b''
            async for chunk in r.aiter_bytes():
                data += chunk
                if len(data) > 10_000_000:
                    raise ValueError('file is over 10 MB')
    fd, p = tempfile.mkstemp(suffix=ext)
    with os.fdopen(fd, 'wb') as fh:
        fh.write(data)
    return p

def n1_frame(df):
    lines = ['rows=%d cols=%d' % (len(df), df.shape[1])]
    for col in list(df.columns)[:30]:
        s = df[col]
        sample = ', '.join(str(v)[:25] for v in s.dropna().unique()[:3])
        lines.append('- %s | %s | null %.0f%% | distinct %d | e.g. %s' % (str(col)[:40], s.dtype, 100.0 * int(s.isna().sum()) / max(len(df), 1), int(s.nunique(dropna=True)), sample))
    return NL.join(lines)

if step_n1_ref:
    try:
        path = await n1_localize(step_n1_ref)
        ext = os.path.splitext(path)[1].lower() or os.path.splitext(urlparse(str(step_n1_ref)).path)[1].lower()
        if ext in ('.csv', '.tsv'):
            step_n1_profile_text = 'CSV profile:' + NL + n1_frame(pd.read_csv(path, sep=None, engine='python'))
        elif ext in ('.xlsx', '.xlsm', '.xls'):
            sheets = pd.read_excel(path, sheet_name=None)
            step_n1_profile_text = NL.join('Sheet ' + str(nm) + ': ' + n1_frame(df) for nm, df in list(sheets.items())[:3])
        elif ext == '.pdf':
            from pypdf import PdfReader
            rd = PdfReader(path)
            txt = NL.join((pg.extract_text() or '') for pg in rd.pages[:6])
            step_n1_profile_text = 'PDF with %d pages. ' % len(rd.pages) + (txt.strip() or 'No extractable text (possibly scanned).')
        elif ext == '.docx':
            from docx import Document
            step_n1_profile_text = 'Word document text: ' + NL.join(p.text for p in Document(path).paragraphs if p.text.strip())
        elif ext in ('.txt', '.md'):
            step_n1_profile_text = open(path, encoding='utf-8', errors='ignore').read()
        else:
            raise ValueError('Unsupported attachment type ' + ext)
        step_n1_profile_text = step_n1_profile_text[:4000]
        step_n1_notices.append('Read the attached sample file.')
    except Exception as exc:
        step_n1_profile_text = ''
        step_n1_notices.append('The sample file could not be used (' + type(exc).__name__ + '). Continuing without it.')

print('Action=%s, turn=%d, known areas=%d, recent turns=%d, attachment chars=%d, state reset=%s' % (step_n1_action, step_n1_state['turn'], len(step_n1_state['fields']), len(step_n1_turns), len(step_n1_profile_text), step_n1_reset))
display(HTML('<div class="ag-root"><h3 class="ag-h3">Conversation loaded</h3><p class="ag-body">Turn ' + str(step_n1_state['turn'] + 1) + '. ' + str(len(step_n1_state['fields'])) + ' of ' + str(len(AREAS)) + ' areas known so far.</p></div>'))
'''

N2 = r'''
from typing import Literal
from pydantic import BaseModel

class N2Update(BaseModel):
    area: Literal['business_goal','current_process','systems_involved','volume_per_day','exception_handling','security_compliance','success_metrics','users_and_roles','data_sources_and_formats','integrations_required','timeline_constraints','stakeholders_approvers']
    value: str
    provenance: Literal['FACT','INFERENCE','ASSUMPTION']
    evidence: str
    specific: bool

class N2Wbs(BaseModel):
    unit: Literal['DISCOVERY_WORKSHOP','PROCESS_MAPPING','INTEGRATION_SIMPLE','INTEGRATION_COMPLEX','DOCUMENT_EXTRACTION_AUTO','WORKFLOW_AUTO_BUILD','DATA_CLEANUP_MIGRATION','POLICY_RULESET','HUMAN_REVIEW_STEP','REPORTING_DASHBOARD','SECURITY_REVIEW','UAT_AND_HYPERCARE','TRAINING_HANDOVER']
    quantity: int
    reason: str

class N2Fit(BaseModel):
    verdict: Literal['BUILD','NOT_YET','DONT_BUILD','UNSURE']
    reasons: list[str]
    kill_risks: list[str]
    ask_to_see: list[str]

class N2Out(BaseModel):
    client_name: str
    updates: list[N2Update]
    wbs_units: list[N2Wbs]
    risks: list[str]
    wants_manual: bool
    acknowledgement: str
    asked_about: Literal['business_goal','current_process','systems_involved','volume_per_day','exception_handling','security_compliance','success_metrics','users_and_roles','data_sources_and_formats','integrations_required','timeline_constraints','stakeholders_approvers']
    next_question: str
    why_asking: str
    proposal_summary: str
    fit: N2Fit

step_n2_known = NL.join('%s: %s [%s]' % (k, step_n1_state['fields'][k]['v'], step_n1_state['fields'][k]['p']) for k, _ in AREAS if k in step_n1_state['fields']) or '(nothing yet)'
step_n2_unknown = ', '.join(k for k, _ in AREAS if k not in step_n1_state['fields']) or '(none)'
step_n2_turns = NL.join(t['role'] + ': ' + t['text'] for t in step_n1_turns) or '(start of conversation)'

step_n2_system = ('You are a digital forward deployed engineer at Supervity. You help a colleague scope a client problem through conversation. Sound like a sharp, friendly senior engineer: plain English, short sentences, no filler, no flattery. Do not use em-dashes. '
 'Each turn: (1) Read the new message and update what is known. Return in updates only the areas this message adds or changes. FACT means the person said it or the attached file shows it, and evidence must then be a verbatim excerpt of at most 25 words copied from the new message, the recent turns or the attachment. specific is false when the answer is vague, such as a lot, soon, some, quickly, or anything without a number, name or concrete detail; a real FDE does not accept vague answers. INFERENCE means you derived it; evidence briefly says why. ASSUMPTION is a sensible default you propose, and the value must say so. Never invent systems, numbers, names or compliance rules. '
 '(2) acknowledgement: one or two sentences saying what you understood, using their words. (3) Ask exactly one question. Choose asked_about as the first area in this priority order that is still not confirmed after your updates: ' + ', '.join(k for k, _ in AREAS) + '. Write next_question about that area, specific to what they told you, one sentence, the way a sharp FDE would ask: concrete, and quoting their words when they were vague. why_asking: one short sentence on why the answer changes what gets built. '
 '(4) proposal_summary: once the goal and systems are known, two to four sentences describing what you would build, in plain terms, as an Auto run by AI Employees with people reviewing exceptions; otherwise an empty string. Make no promises about savings or timelines. '
 '(5) wants_manual is true only if the person asks for the manual, a document or a write-up. (6) wbs_units: the full current list of work units that fit the scope as understood, counted as whole numbers; empty if the scope is still unclear. risks: up to 4 short items. client_name: the client if mentioned, else an empty string. '
 '(7) fit: judge like a seasoned FDE whether this should be built at all, and be willing to say no. DONT_BUILD if the work is too rare, too variable, or cheaper to leave manual. NOT_YET if the process is not defined, the data is not reachable, or no one owns the outcome. BUILD only if volume, a stable process, an owner and data access all support it. UNSURE while the goal, process, volume or systems are unknown. reasons: up to 3 plain reasons tied to the stated facts. kill_risks: up to 3 people or functions who could block or kill the project, such as IT security, legal, finance or a team lead whose work changes, with a few words on why; only from what the facts suggest. ask_to_see: up to 3 concrete things to look at first, such as 10 real invoices from last week or the shared mailbox. Leave all three lists empty when the verdict is UNSURE and little is known.')
step_n2_user = ('KNOWN SO FAR:' + NL + step_n2_known + NL + 'NOT YET KNOWN: ' + step_n2_unknown + NL + 'CLIENT NAME: ' + (step_n1_state['client_name'] or '(unknown)') + NL + 'RECENT TURNS:' + NL + step_n2_turns + NL + 'ATTACHMENT PROFILE:' + NL + (step_n1_profile_text or '(none)') + NL + 'NEW MESSAGE:' + NL + step_n1_msg)

step_n2_res = await ai_generate(system_prompt=step_n2_system, user_prompt=step_n2_user, response_schema_pydantic_model=N2Out, model='everyday', thinking_level='low')
print('Model returned %d updates, %d work units, asked about %s.' % (len(step_n2_res.updates), len(step_n2_res.wbs_units), step_n2_res.asked_about))
'''

N3 = r'''
import re, html, json
from IPython.display import display, HTML

def n3_norm(s):
    return re.sub(r'\s+', ' ', str(s or '')).lower().strip()

N3_UNITS = {'zero':0,'one':1,'two':2,'three':3,'four':4,'five':5,'six':6,'seven':7,'eight':8,'nine':9,'ten':10,'eleven':11,'twelve':12,'thirteen':13,'fourteen':14,'fifteen':15,'sixteen':16,'seventeen':17,'eighteen':18,'nineteen':19}
N3_TENS = {'twenty':20,'thirty':30,'forty':40,'fifty':50,'sixty':60,'seventy':70,'eighty':80,'ninety':90}

def n3_tokens(s):
    t = str(s or '').lower().replace(',', '')
    t = re.sub(r'(\d)\s*k\b', lambda m: m.group(1) + '000', t)
    t = re.sub(r'[^a-z0-9.%]+', ' ', t)
    raw = [x.strip('.') for x in t.split()]
    raw = [x for x in raw if x]
    out, i = [], 0
    while i < len(raw):
        w = raw[i]
        if w in N3_UNITS or w in N3_TENS or w in ('hundred', 'thousand'):
            total, cur, used = 0, 0, False
            while i < len(raw) and (raw[i] in N3_UNITS or raw[i] in N3_TENS or raw[i] in ('hundred', 'thousand', 'and')):
                x = raw[i]
                if x == 'and':
                    if i + 1 < len(raw) and (raw[i + 1] in N3_UNITS or raw[i + 1] in N3_TENS):
                        i += 1
                        continue
                    break
                if x in N3_UNITS: cur += N3_UNITS[x]
                elif x in N3_TENS: cur += N3_TENS[x]
                elif x == 'hundred': cur = max(cur, 1) * 100
                elif x == 'thousand': total += max(cur, 1) * 1000; cur = 0
                used = True
                i += 1
            out.append(str(total + cur) if used else w)
        else:
            out.append(w)
            i += 1
    return out

def n3_matches(evidence, source_tokens):
    et = n3_tokens(evidence)
    if len(et) < 2:
        return False
    sj = ' ' + ' '.join(source_tokens) + ' '
    if ' ' + ' '.join(et) + ' ' in sj:
        return True
    import difflib
    sm = difflib.SequenceMatcher(None, et, source_tokens, autojunk=False)
    matched = sum(b.size for b in sm.get_matching_blocks())
    longest = max((b.size for b in sm.get_matching_blocks()), default=0)
    return matched >= 0.9 * len(et) and longest >= max(2, int(0.5 * len(et)))

def n3_scrub(s):
    s = str(s or '').replace(chr(8212), ', ').replace(chr(8211), '-')
    return re.sub(r'\s+', ' ', s).strip()

STRENGTH = {'FACT': 3, 'INFERENCE': 2, 'ASSUMPTION': 1}
step_n3_source = n3_tokens(step_n1_msg + ' ' + ' '.join(t['text'] for t in step_n1_turns) + ' ' + step_n1_profile_text)
step_n3_fields = dict(step_n1_state['fields'])
step_n3_downgraded = []
for u in step_n2_res.updates:
    v = u.value.strip()[:MAXV]
    if not v or u.area not in LABELS:
        continue
    p = u.provenance
    e = u.evidence.strip()[:MAXE]
    if not u.specific:
        p, e = 'ASSUMPTION', 'vague'
    if p == 'FACT':
        if not n3_matches(e, step_n3_source):
            p = 'INFERENCE'
            step_n3_downgraded.append(LABELS[u.area])
    old = step_n3_fields.get(u.area)
    if old and STRENGTH[old['p']] > STRENGTH[p]:
        continue
    step_n3_fields[u.area] = {'v': v, 'p': p, 'e': e}

step_n3_wbs = list(step_n1_state['wbs'])
step_n3_new = {}
for w in step_n2_res.wbs_units:
    if w.unit in RATE:
        step_n3_new[w.unit] = {'unit': w.unit, 'quantity': min(max(int(w.quantity), 1), 50), 'reason': n3_scrub(w.reason)[:200]}
if step_n3_new:
    step_n3_wbs = list(step_n3_new.values())
step_n3_client = step_n1_state['client_name'] or n3_scrub(step_n2_res.client_name)[:200]
step_n3_risks = [n3_scrub(r)[:200] for r in step_n2_res.risks][:6] or step_n1_state['risks']

def n3_ok(k):
    return k in step_n3_fields and step_n3_fields[k]['p'] in ('FACT', 'INFERENCE')

step_n3_total_w = sum(2 if k in CRITICAL else 1 for k, _ in AREAS)
step_n3_ok_w = sum(2 if k in CRITICAL else 1 for k, _ in AREAS if n3_ok(k))
step_n3_completeness = round(step_n3_ok_w / step_n3_total_w, 2)
step_n3_gaps = []
for k, label in AREAS:
    if n3_ok(k):
        continue
    status = ('VAGUE' if step_n3_fields[k]['e'] == 'vague' else 'ASSUMED') if k in step_n3_fields else 'MISSING'
    if status == 'VAGUE':
        q = 'You said "' + step_n3_fields[k]['v'][:120] + '", which is too vague to build from. ' + QUESTIONS[k]
    elif status == 'ASSUMED':
        q = 'Please confirm this assumption: ' + step_n3_fields[k]['v'][:150] + ' ' + QUESTIONS[k]
    else:
        q = QUESTIONS[k]
    step_n3_gaps.append({'key': k, 'label': label, 'status': status, 'critical': k in CRITICAL, 'question': q})
step_n3_critical_open = [g for g in step_n3_gaps if g['critical']]
step_n3_key_total = len(CRITICAL)
step_n3_key_confirmed = sum(1 for k in CRITICAL if n3_ok(k))

mf = step_n2_res.fit
step_n3_fit = {'verdict': mf.verdict, 'reasons': [n3_scrub(x)[:200] for x in mf.reasons][:3], 'kill_risks': [n3_scrub(x)[:200] for x in mf.kill_risks][:3], 'ask_to_see': [n3_scrub(x)[:160] for x in mf.ask_to_see][:3]}
if not (n3_ok('business_goal') and n3_ok('volume_per_day') and n3_ok('current_process')):
    step_n3_fit['verdict'] = 'UNSURE'
step_n3_low_volume = False
_vol = step_n3_fields.get('volume_per_day')
if _vol and _vol['p'] == 'FACT':
    _nums = [float(x) for x in re.findall(r'\d+(?:\.\d+)?', ' '.join(n3_tokens(_vol['v'])))]
    if _nums and max(_nums) < 20:
        step_n3_low_volume = True
if step_n3_low_volume and step_n3_fit['verdict'] == 'BUILD':
    step_n3_fit['verdict'] = 'UNSURE'
    step_n3_fit['reasons'] = (['Volume is under 20 items a day, so the build may not pay back.'] + step_n3_fit['reasons'])[:3]
step_n3_readiness = 'NEEDS_CLARIFICATION' if step_n3_critical_open else 'READY_FOR_REVIEW'

step_n3_advise = step_n3_fit['verdict'] in ('NOT_YET', 'DONT_BUILD')
step_n3_wants = step_n1_action == 'manual' or (bool(step_n2_res.wants_manual) and re.search(r'manual|document|write.?up|brief', step_n1_msg, re.I) is not None)
if step_n3_wants:
    step_n3_phase = 'manual'
elif step_n3_advise:
    step_n3_phase = 'advise'
elif step_n3_critical_open:
    step_n3_phase = 'discover'
else:
    step_n3_phase = 'propose'

step_n3_ack = n3_scrub(step_n2_res.acknowledgement) or 'Got it.'
if step_n3_phase == 'discover':
    g = step_n3_gaps[0]
    nq = n3_scrub(step_n2_res.next_question)
    if g['status'] == 'MISSING' and step_n2_res.asked_about == g['key'] and 10 <= len(nq) <= 400 and nq.endswith('?'):
        question = nq
    else:
        question = g['question']
    why = n3_scrub(step_n2_res.why_asking)
    if g['status'] == 'MISSING' and 10 <= len(why) <= 220:
        question = question + ' (' + why.rstrip('.') + '.)'
    if step_n1_state['turn'] % 3 == 2 and not step_n1_profile_text and step_n3_fit['ask_to_see']:
        question = question + ' It would also help to see ' + step_n3_fit['ask_to_see'][0].rstrip('.') + '.'
    step_n3_reply = step_n3_ack + ' ' + question
elif step_n3_phase == 'advise':
    head = 'I would not build this at all.' if step_n3_fit['verdict'] == 'DONT_BUILD' else 'I would not start building yet.'
    why_txt = ' '.join(r.rstrip('.') + '.' for r in step_n3_fit['reasons'])
    risk_txt = (' The people who could stop it: ' + '; '.join(step_n3_fit['kill_risks']) + '.') if step_n3_fit['kill_risks'] else ''
    see_txt = (' First I would look at ' + step_n3_fit['ask_to_see'][0].rstrip('.') + '.') if step_n3_fit['ask_to_see'] else ''
    step_n3_reply = step_n3_ack + ' To be straight with you: ' + head + ' ' + why_txt + risk_txt + see_txt + ' Tell me if I have misread anything, or say create the manual and I will write up the recommendation.'
elif step_n3_phase == 'propose':
    prop = n3_scrub(step_n2_res.proposal_summary)
    if not prop:
        prop = 'Based on what we have so far, I would automate this: ' + step_n3_fields['business_goal']['v'] + '.'
    rest = len(step_n3_gaps)
    tail = (' There are ' + str(rest) + ' smaller areas still open, and the manual will list them.') if rest else ''
    risk_txt = (' One thing that could stop it: ' + step_n3_fit['kill_risks'][0].rstrip('.') + '.') if step_n3_fit['kill_risks'] else ''
    see_txt = (' Before we build, I would want to see ' + step_n3_fit['ask_to_see'][0].rstrip('.') + '.') if step_n3_fit['ask_to_see'] else ''
    step_n3_reply = step_n3_ack + ' Here is what I would propose. ' + prop + tail + risk_txt + see_txt + ' Say create the manual when you are ready, or tell me what to change.'
else:
    step_n3_reply = step_n3_ack

step_n3_lines = []
step_n3_base = 0
for w in step_n3_wbs:
    hrs = RATE[w['unit']] * w['quantity']
    step_n3_base += hrs
    step_n3_lines.append({'unit': w['unit'], 'quantity': w['quantity'], 'hours_each': RATE[w['unit']], 'hours': hrs, 'reason': w['reason']})
step_n3_high = round(step_n3_base * (1.25 + 0.5 * (1 - step_n3_completeness)))
step_n3_estimate = {'lines': step_n3_lines, 'low_hours': step_n3_base, 'high_hours': step_n3_high, 'low_days': round(step_n3_base / 8, 1), 'high_days': round(step_n3_high / 8, 1), 'calibrated': RATE_SOURCE == 'supplied', 'basis': ('YOUR RATE CARD' if RATE_SOURCE == 'supplied' else 'STARTER RATE CARD, not calibrated to delivery history')}

step_n3_state = {'client_name': step_n3_client, 'fields': step_n3_fields, 'wbs': step_n3_wbs, 'risks': step_n3_risks, 'turn': step_n1_state['turn'] + 1}
while len(json.dumps(step_n3_state)) > 11000:
    for f in step_n3_state['fields'].values():
        f['e'] = f['e'][:60]
        f['v'] = f['v'][:250]
    if len(json.dumps(step_n3_state)) > 11000:
        step_n3_state['risks'] = []
        step_n3_state['wbs'] = step_n3_state['wbs'][:8]
        break

print('Key areas confirmed %d of %d. Fit=%s. Completeness %.0f%%. Phase=%s. Open areas=%d (critical %d). Unverified FACT downgraded: %s.' % (step_n3_key_confirmed, step_n3_key_total, step_n3_fit['verdict'], 100 * step_n3_completeness, step_n3_phase, len(step_n3_gaps), len(step_n3_critical_open), ', '.join(step_n3_downgraded) or 'none'))
display(HTML('<div class="ag-root"><h3 class="ag-h3">Understanding updated</h3><p class="ag-body">Completeness ' + str(int(100 * step_n3_completeness)) + '%. Next: ' + html.escape(step_n3_phase) + '.</p></div>'))
'''

N4 = r'''
import os, re, json, html, tempfile
from datetime import datetime
from zoneinfo import ZoneInfo
from pydantic import BaseModel
from IPython.display import display, HTML

step_n4_manual = None
if step_n3_phase != 'manual':
    print('Manual not requested this turn.')
else:
    class N4BA(BaseModel):
        area: str
        today: str
        with_the_auto: str
    class N4Step(BaseModel):
        title: str
        description: str
        owner: str
    class N4Role(BaseModel):
        role: str
        responsibility: str
    class N4Exc(BaseModel):
        situation: str
        handling: str
    class N4Phase(BaseModel):
        name: str
        outcome: str
        activities: list[str]
    class N4Bp(BaseModel):
        step_name: str
        what_it_does: str
        worker: str
        human_checkpoint: bool
    class N4Doc(BaseModel):
        title: str
        recommendation: str
        kill_risks: list[str]
        see_first: list[str]
        use_case_line: str
        summary: str
        problem: list[str]
        proposed_solution: str
        before_and_after: list[N4BA]
        how_it_works: list[N4Step]
        blueprint_steps: list[N4Bp]
        blueprint_inputs: list[str]
        blueprint_policies: list[str]
        team_inputs: list[str]
        roles: list[N4Role]
        exceptions: list[N4Exc]
        security: list[str]
        rollout: list[N4Phase]
        success_measures: list[str]
        risks: list[str]
        next_step: str

    BAN = {'seamlessly': 'smoothly', 'seamless': 'smooth', 'leveraging': 'using', 'leverage': 'use', 'robust': 'reliable', 'world-class': 'strong', 'cutting-edge': 'modern', 'unlock': 'enable', 'supercharge': 'speed up'}
    n4_hits = [0]
    def n4_scrub(s):
        s0 = n3_scrub(s)
        s1 = re.sub(r'AI agents', 'AI Employees', s0)
        s1 = re.sub(r'AI agent', 'AI Employee', s1)
        s1 = re.sub(r'AI assistants', 'AI Employees', s1)
        s1 = re.sub(r'AI assistant', 'AI Employee', s1)
        for w, r in BAN.items():
            s1 = re.sub(r'\b' + w + r'\b', r, s1, flags=re.I)
        if s1 != s0:
            n4_hits[0] += 1
        return s1

    facts = {k: {'v': f['v'], 'basis': f['p']} for k, f in step_n3_fields.items()}
    payload = json.dumps({'client': step_n3_client or 'the client', 'known': facts, 'not_known': [LABELS[g['key']] for g in step_n3_gaps], 'work_units': [{'unit': l['unit'], 'qty': l['quantity'], 'why': l['reason']} for l in step_n3_lines], 'risks_seen': step_n3_risks, 'build_verdict': step_n3_fit['verdict'], 'verdict_reasons': step_n3_fit['reasons'], 'possible_blockers': step_n3_fit['kill_risks'], 'look_at_first': step_n3_fit['ask_to_see']}, separators=(',', ':'))
    system = ('You write a solution manual for a Supervity delivery team and the client stakeholders. It describes the one solution being proposed, using only the facts given. Write in natural executive English: plain, direct sentences a senior professional would say to a colleague. State points positively. Do not use em-dashes, slogans, rhetorical closers or filler such as seamless, leverage, robust or world-class. '
              'Use the product nouns correctly: the solution is an Auto, and the workers are AI Employees. Name Auto Policies only for rules, and the Auto Workbench only where people review exceptions. Never write AI agent, AI tool or AI assistant. '
              'Never invent numbers, customer names, systems, dates, savings or guarantees, and make no ROI promise. Where something is not known, write To be confirmed with the client. Mark statements resting on an inference with (inferred) and on an assumption with (assumption). '
              'Keep it concrete and specific to this client function. summary under 100 words. how_it_works has 4 to 8 steps in the order the work happens, each description under 40 words, each with an owner that is an AI Employee or a named role. before_and_after has 3 to 5 rows. rollout has 3 phases. exceptions has 2 to 5 rows drawn from the facts or marked To be confirmed. next_step is one clear action. recommendation is a short paragraph on whether to build this, stating the given build_verdict (BUILD, NOT_YET, DONT_BUILD or UNSURE) and its reasons without changing the verdict; when the verdict is NOT_YET or DONT_BUILD, the solution sections describe what would have to be true before building. kill_risks lists who could stop the project and why, starting from possible_blockers. see_first lists what to look at before building, starting from look_at_first. blueprint_steps is the build plan for the Auto in execution order, 4 to 8 steps, each with a worker that is an AI Employee or code, and human_checkpoint true where a person approves; blueprint_inputs lists what a run needs as input; blueprint_policies lists business rules that belong in Auto Policies, marked To be confirmed where unknown.')
    doc_res = await ai_generate(system_prompt=system, user_prompt=payload, response_schema_pydantic_model=N4Doc, model='everyday', thinking_level='low')

    from docx import Document
    from docx.shared import RGBColor
    d = Document()
    for hn in ('Heading 1', 'Heading 2', 'Title'):
        d.styles[hn].font.color.rgb = RGBColor(0x8C, 0x0F, 0x19)
    def hd(t):
        d.add_heading(t, 1)
    def para(t):
        d.add_paragraph(n4_scrub(t))
    def bullets(items):
        for it in items:
            d.add_paragraph(n4_scrub(it), style='List Bullet')
    def table(headers, rows):
        t = d.add_table(rows=1, cols=len(headers))
        t.style = 'Table Grid'
        for i, h in enumerate(headers):
            t.rows[0].cells[i].text = h
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = n4_scrub(v)

    step_n4_title = n4_scrub(doc_res.title) or ('Solution manual: ' + (step_n3_client or 'client'))
    d.add_heading(step_n4_title, 0)
    now = datetime.now(ZoneInfo('Asia/Kolkata')).strftime('%d %b %Y %H:%M IST')
    d.add_paragraph('Draft prepared by the digital FDE on ' + now + '. A person must review it before it is used. Completeness: ' + str(int(100 * step_n3_completeness)) + '%.')
    hd('The use case'); para(doc_res.use_case_line)
    VLAB = {'BUILD': 'Build', 'NOT_YET': 'Not yet', 'DONT_BUILD': 'Do not build', 'UNSURE': 'Not enough information to say'}
    hd('Recommendation')
    d.add_paragraph('Verdict: ' + VLAB[step_n3_fit['verdict']])
    para(doc_res.recommendation)
    hd('Who could stop this'); bullets(doc_res.kill_risks or ['To be confirmed with the client'])
    hd('What to look at before building'); bullets(doc_res.see_first or ['To be confirmed with the client'])
    hd('Summary'); para(doc_res.summary)
    hd('The problem'); bullets(doc_res.problem)
    hd('What we propose'); para(doc_res.proposed_solution)
    hd('Before and after'); table(['Area', 'Today', 'With the Auto'], [[b.area, b.today, b.with_the_auto] for b in doc_res.before_and_after])
    hd('How it works')
    for s in doc_res.how_it_works:
        d.add_paragraph(n4_scrub(s.title) + '. ' + n4_scrub(s.description) + ' (' + n4_scrub(s.owner) + ')', style='List Number')
    hd('Proposed Auto blueprint')
    para('This is the plan for building the Auto. It is not built yet.')
    table(['Step', 'What it does', 'Worker', 'Person approves'], [[b.step_name, b.what_it_does, b.worker, 'Yes' if b.human_checkpoint else 'No'] for b in doc_res.blueprint_steps])
    para('Inputs a run needs:')
    bullets(doc_res.blueprint_inputs)
    para('Rules to set up as Auto Policies:')
    bullets(doc_res.blueprint_policies)
    hd('What the team needs to provide'); bullets(doc_res.team_inputs)
    hd('Roles'); table(['Role', 'Responsibility'], [[r.role, r.responsibility] for r in doc_res.roles])
    hd('Exceptions and who handles them'); table(['Situation', 'Handling'], [[e.situation, e.handling] for e in doc_res.exceptions])
    hd('Security and compliance'); bullets(doc_res.security)
    hd('Rollout')
    for ph in doc_res.rollout:
        d.add_heading(n4_scrub(ph.name), 2)
        para(ph.outcome)
        bullets(ph.activities)
    hd('How success is measured'); bullets(doc_res.success_measures)
    hd('Risks'); bullets(list(doc_res.risks) + step_n3_risks)
    hd('Open questions for the client')
    if step_n3_gaps:
        for g in step_n3_gaps:
            d.add_paragraph(('Critical. ' if g['critical'] else '') + g['question'], style='List Number')
    else:
        para('No open items.')
    hd('What we know and how we know it')
    table(['Area', 'What we have', 'Basis'], [[LABELS[k], f['v'], f['p']] for k, f in step_n3_fields.items()])
    d.add_paragraph('FACT: stated by the person or shown in an attached file. INFERENCE: derived from what was said. ASSUMPTION: a proposed default that needs confirmation.')
    if step_n3_lines:
        hd('Indicative effort')
        d.add_paragraph('Range: ' + str(step_n3_estimate['low_hours']) + ' to ' + str(step_n3_estimate['high_hours']) + ' hours. Basis: ' + step_n3_estimate['basis'] + '. The upper end widens when fewer areas are confirmed.')
        table(['Work unit', 'Qty', 'Hours', 'Why'], [[l['unit'], str(l['quantity']), str(l['hours']), l['reason']] for l in step_n3_lines])
    hd('Next step'); para(doc_res.next_step)
    hd('Review')
    d.add_paragraph('Reviewer: ______________________    Date: ______________')
    d.add_paragraph('Decision:  [ ] Approved    [ ] Needs changes')

    out_dir = tempfile.mkdtemp()
    safe = re.sub(r'[^A-Za-z0-9_-]+', '_', step_n3_client)[:40] or 'client'
    path = os.path.join(out_dir, 'Solution_Manual_' + safe + '.docx')
    d.save(path)
    url = await upload_output_file(path)
    step_n4_manual = {'url': url, 'title': step_n4_title, 'verdict': step_n3_fit['verdict'], 'blueprint': [{'step': n4_scrub(b.step_name), 'does': n4_scrub(b.what_it_does), 'worker': n4_scrub(b.worker), 'review': bool(b.human_checkpoint)} for b in doc_res.blueprint_steps][:8], 'open_questions': len(step_n3_gaps), 'wording_fixes': n4_hits[0]}
    print('Manual written and uploaded. Open questions: %d. Wording fixes applied: %d.' % (len(step_n3_gaps), n4_hits[0]))
    display(HTML('<div class="ag-root"><h3 class="ag-h3">Manual ready</h3><p class="ag-body"><a class="ag-link" href="' + html.escape(url) + '">Download the manual</a></p></div>'))
'''

N5 = r'''
import json, html
from IPython.display import display, HTML

step_n5_areas = []
for k, label in AREAS:
    f = step_n3_fields.get(k)
    step_n5_areas.append({'key': k, 'label': label, 'status': f['p'] if f else 'UNKNOWN', 'value': f['v'] if f else '', 'critical': k in CRITICAL})

step_n5_reply = step_n3_reply
if step_n3_phase == 'manual':
    if step_n4_manual:
        n = step_n4_manual['open_questions']
        step_n5_reply = 'The manual is ready, with a blueprint of the Auto I would build. ' + (('It lists ' + str(n) + ' open questions the client still needs to answer.') if n else 'Every area is covered.') + ' Tell me what to change and I will update it.'
    else:
        step_n5_reply = 'I could not produce the manual this time. Please try again.'

step_n5_resp = {
    'reply': step_n5_reply,
    'phase': step_n3_phase,
    'agent_version': AGENT_VERSION,
    'key_confirmed': step_n3_key_confirmed,
    'key_total': step_n3_key_total,
    'fit': step_n3_fit,
    'completeness': step_n3_completeness,
    'readiness': step_n3_readiness,
    'areas': step_n5_areas,
    'open_questions': [{'label': g['label'], 'question': g['question'], 'critical': g['critical']} for g in step_n3_gaps][:6],
    'estimate': ({'low_hours': step_n3_estimate['low_hours'], 'high_hours': step_n3_estimate['high_hours'], 'calibrated': step_n3_estimate['calibrated'], 'basis': step_n3_estimate['basis']} if step_n3_lines else None),
    'manual': ({'url': step_n4_manual['url'], 'title': step_n4_manual['title'], 'verdict': step_n4_manual['verdict'], 'blueprint': step_n4_manual['blueprint']} if step_n4_manual else None),
    'notices': step_n1_notices,
    'state': step_n3_state,
}
print(json.dumps({'fde_response': step_n5_resp}, indent=2, ensure_ascii=False))
display(HTML('<div class="ag-root"><h3 class="ag-h3">Reply</h3><p class="ag-body">' + html.escape(step_n5_reply) + '</p></div>'))
'''

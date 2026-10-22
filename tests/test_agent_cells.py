"""Runs the Digital FDE agent's step code (agent/cells.py, the exact code published to the
workflow) against a scripted model. This covers everything that is plain code: fact
verification, gap targeting, fit rules, state handling, manual building, output format.
It cannot judge the real model's wording; only a live run can."""
import ast
import asyncio
import contextlib
import io
import json
import os
import tempfile

import IPython.display as _d
import pytest

import re

from agent import cells as C

_d.display = lambda *a, **k: None
AGENT_VERSION = re.search(r"AGENT_VERSION = '([^']+)'", C.N1).group(1)
RATE_KEYS = re.findall(r"'([A-Z_]+)':\d+", re.search(r"RATE = \{.*?\}", C.N1).group(0))
FLAGS = ast.PyCF_ALLOW_TOP_LEVEL_AWAIT
ALL = [C.N1, C.N2, C.N3, C.N4, C.N5]


def make(ui, n2, n4=None, calls=None):
    async def ai_generate(system_prompt, user_prompt, response_schema_pydantic_model, model="everyday", **kw):
        nm = response_schema_pydantic_model.__name__
        if calls is not None:
            calls.append((nm, len(system_prompt) + len(user_prompt)))
        return response_schema_pydantic_model.model_validate(n2 if nm == "N2Out" else n4)

    async def upload_output_file(p):
        assert os.path.exists(p) and os.path.getsize(p) > 5000
        return "https://files.test/" + os.path.basename(p)

    return {"user_inputs": ui, "ai_generate": ai_generate, "upload_output_file": upload_output_file, "__name__": "cell"}


async def run(src, g):
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        out = eval(compile(src, "cell", "exec", flags=FLAGS), g)
        if asyncio.iscoroutine(out):
            await out
    return buf.getvalue()


async def turn(ui, n2, n4=None, calls=None, upto=5):
    g = make(ui, n2, n4, calls)
    outs = [await run(s, g) for s in ALL[:upto]]
    return g, outs


def resp(outs):
    return json.loads(outs[4])["fde_response"]


def area(r, key):
    return [a for a in r["areas"] if a["key"] == key][0]


FIT0 = dict(verdict="UNSURE", reasons=[], kill_risks=[], ask_to_see=[])


def n2d(**o):
    d = dict(client_name="", updates=[], wbs_units=[], risks=[], wants_manual=False, acknowledgement="Got it.",
             asked_about="current_process", next_question="Can you walk me through how this gets done today?",
             why_asking="", proposal_summary="", fit=FIT0)
    d.update(o)
    return d


def up(area_, value, ev, prov="FACT", specific=True):
    return dict(area=area_, value=value, provenance=prov, evidence=ev, specific=specific)


MSG = "We want to automate our invoice thing, it takes too long. Finance handles it in SAP and a lot comes by email as PDFs. Maybe 300 a day."
FULL = [up("business_goal", "Automate invoice handling", "automate our invoice thing"),
        up("systems_involved", "SAP, email", "Finance handles it in SAP"),
        up("volume_per_day", "About 300 a day", "Maybe 300 a day")]
KEYS = ["business_goal", "current_process", "systems_involved", "volume_per_day", "exception_handling", "security_compliance", "success_metrics", "stakeholders_approvers"]
STATE = {"client_name": "Acme", "fields": {k: {"v": "v" + k[:3], "p": "FACT", "e": ""} for k in KEYS}, "wbs": [], "risks": [], "turn": 4}

N4 = dict(title="Solution manual: Acme", recommendation="Verdict: build. Volume supports it.", kill_risks=["IT security: data leaves SAP"], see_first=["10 real invoices"],
          use_case_line="Invoice processing for Acme finance.", summary="Acme will run invoice intake as an Auto. AI agents read each invoice.",
          problem=["Invoices take too long, we leverage manual effort."], proposed_solution="An Auto reads invoices and posts them to SAP — seamlessly.",
          before_and_after=[dict(area="Intake", today="Manual", with_the_auto="Automatic")] * 3,
          how_it_works=[dict(title="Read", description="The AI assistant reads the PDF.", owner="Invoice AI Employee")] * 4,
          blueprint_steps=[dict(step_name="Read invoices", what_it_does="Reads PDFs", worker="Invoice AI Employee", human_checkpoint=False)] * 4,
          blueprint_inputs=["Invoice files"], blueprint_policies=["Approval limits: To be confirmed with the client"],
          team_inputs=["Sample invoices"], roles=[dict(role="AP lead", responsibility="Reviews exceptions")],
          exceptions=[dict(situation="Mismatch", handling="To be confirmed with the client")] * 2, security=["To be confirmed with the client"],
          rollout=[dict(name="Discover", outcome="Scope agreed", activities=["Workshop"])] * 3, success_measures=["Cycle time"], risks=["Data quality"], next_step="Hold a discovery workshop.")


async def test_first_turn_shape_and_cost():
    calls = []
    g, outs = await turn({"user_message": MSG}, n2d(client_name="Acme", updates=FULL, acknowledgement="You want invoices automated — understood."), calls=calls)
    r = resp(outs)
    assert outs[4].lstrip().startswith("{")           # only JSON on stdout
    assert len(calls) == 1                            # one model call per chat turn
    assert r["phase"] == "discover" and r["agent_version"] == AGENT_VERSION
    assert "—" not in r["reply"] and r["reply"].count("?") == 1
    assert area(r, "business_goal")["status"] == "FACT" and r["manual"] is None
    assert r["state"]["turn"] == 1 and len(json.dumps(r["state"])) < 12000
    assert r["key_total"] == 8 and r["key_confirmed"] == 3


async def test_fabricated_fact_is_downgraded():
    _, outs = await turn({"user_message": MSG}, n2d(updates=[up("security_compliance", "SOC2 required", "we need SOC2 compliance")]))
    assert area(resp(outs), "security_compliance")["status"] == "INFERENCE"


@pytest.mark.parametrize("said,evidence", [
    ("About three hundred invoices a day", "300 invoices/day"),
    ("We process about 5k invoices a month", "5000 invoices a month"),
    ("it is roughly three hundred and fifty a day", "350 a day"),
    ("We use SAP, Outlook and a shared mailbox.", "SAP Outlook and a shared mailbox"),
])
async def test_fact_match_tolerates_numbers_and_punctuation(said, evidence):
    _, outs = await turn({"user_message": said}, n2d(updates=[up("volume_per_day", "x", evidence)]))
    assert area(resp(outs), "volume_per_day")["status"] == "FACT"


@pytest.mark.parametrize("said,evidence", [
    ("About three hundred invoices a day", "3000 invoices a day"),
    ("About three hundred invoices a day", "about thirty invoices a day"),
    ("We use SAP", "we use Oracle for everything"),
])
async def test_fact_match_still_rejects_wrong_numbers_and_inventions(said, evidence):
    _, outs = await turn({"user_message": said}, n2d(updates=[up("volume_per_day", "x", evidence)]))
    assert area(resp(outs), "volume_per_day")["status"] == "INFERENCE"


async def test_gap_targeting_ignores_models_pick():
    _, outs = await turn({"user_message": MSG}, n2d(updates=FULL, asked_about="timeline_constraints", next_question="What is the deadline for this project?"))
    assert "walk me through" in resp(outs)["reply"]


@pytest.mark.parametrize("bad", ["{not json", "[1,2]", "x" * 15000])
async def test_malformed_state_resets_with_notice(bad):
    _, outs = await turn({"user_message": MSG, "state_json": bad}, n2d())
    r = resp(outs)
    assert r["state"]["turn"] == 1 and not r["state"]["fields"] and any("fresh" in n for n in r["notices"])


async def test_junk_fields_in_state_are_dropped():
    bad = '{"fields":{"business_goal":{"v":"x","p":"BOGUS"}},"wbs":[{"unit":"NOPE","quantity":"a"}],"turn":"zz"}'
    _, outs = await turn({"user_message": MSG, "state_json": bad}, n2d())
    r = resp(outs)
    assert r["state"]["turn"] == 1 and not r["state"]["wbs"] and not r["state"]["fields"]


async def test_empty_message_rejected_and_long_message_capped():
    with pytest.raises(ValueError, match="empty"):
        await run(C.N1, make({"user_message": "   "}, n2d()))
    g = make({"user_message": "x" * 50000}, n2d())
    await run(C.N1, g)
    assert len(g["step_n1_msg"]) == 2000


async def test_propose_phase_and_exact_hours():
    n2 = n2d(proposal_summary="An Auto reads invoices and posts to SAP.",
             fit=dict(verdict="BUILD", reasons=["Volume supports it"], kill_risks=["IT security"], ask_to_see=["10 real invoices"]),
             wbs_units=[dict(unit="WORKFLOW_AUTO_BUILD", quantity=1, reason="build"), dict(unit="INTEGRATION_SIMPLE", quantity=2, reason="SAP")],
             updates=[up("volume_per_day", "300 a day", "300 a day")])
    _, outs = await turn({"user_message": "300 a day, ok", "state_json": json.dumps(STATE)}, n2)
    r = resp(outs)
    assert r["phase"] == "propose" and "create the manual" in r["reply"].lower()
    assert "IT security" in r["reply"] and "10 real invoices" in r["reply"]
    assert r["estimate"]["low_hours"] == 96 and r["estimate"]["high_hours"] > 96
    assert r["estimate"]["calibrated"] is False and "STARTER" in r["estimate"]["basis"]
    assert r["readiness"] == "READY_FOR_REVIEW" and r["key_confirmed"] == 8


async def test_agent_can_say_do_not_build():
    n2 = n2d(updates=[up("business_goal", "Automate a monthly report", "automate our report"), up("volume_per_day", "Twice a month", "twice a month"),
                      up("current_process", "One person does it in Excel", "one person does it in Excel")],
             fit=dict(verdict="DONT_BUILD", reasons=["The work happens twice a month"], kill_risks=["The analyst whose job changes"], ask_to_see=["last month's report"]))
    _, outs = await turn({"user_message": "We want to automate our report. It runs twice a month and one person does it in Excel."}, n2)
    r = resp(outs)
    assert r["phase"] == "advise" and "would not build this at all" in r["reply"] and "analyst" in r["reply"]
    assert r["fit"]["verdict"] == "DONT_BUILD"


async def test_verdict_ignored_until_goal_process_volume_known():
    n2 = n2d(fit=dict(verdict="DONT_BUILD", reasons=["Guess"], kill_risks=[], ask_to_see=[]))
    _, outs = await turn({"user_message": "hi there"}, n2)
    assert resp(outs)["fit"]["verdict"] == "UNSURE" and resp(outs)["phase"] == "discover"


async def test_low_volume_blocks_a_build_verdict():
    st = json.loads(json.dumps(STATE))
    st["fields"]["volume_per_day"] = {"v": "5 a day", "p": "FACT", "e": ""}
    n2 = n2d(fit=dict(verdict="BUILD", reasons=["Looks fine"], kill_risks=[], ask_to_see=[]))
    _, outs = await turn({"user_message": "ok", "state_json": json.dumps(st)}, n2)
    r = resp(outs)
    assert r["fit"]["verdict"] == "UNSURE" and "under 20" in r["fit"]["reasons"][0]


async def test_sample_request_appears_on_third_turn_only_without_attachment():
    st = {"client_name": "", "fields": {}, "wbs": [], "risks": [], "turn": 2}
    fit = dict(verdict="UNSURE", reasons=[], kill_risks=[], ask_to_see=["10 real invoices from last week"])
    _, outs = await turn({"user_message": MSG, "state_json": json.dumps(st)}, n2d(fit=fit))
    assert "10 real invoices from last week" in resp(outs)["reply"]
    st["turn"] = 3
    _, outs = await turn({"user_message": MSG, "state_json": json.dumps(st)}, n2d(fit=fit))
    assert "10 real invoices" not in resp(outs)["reply"]


async def _manual_text(state, n2=None):
    from docx import Document
    box = {}
    g = make({"action": "manual", "user_message": "Create the manual", "state_json": json.dumps(state)}, n2 or n2d(), N4)
    orig = g["upload_output_file"]

    async def cap(p):
        d = Document(p)
        box["txt"] = "\n".join(x.text for x in d.paragraphs) + "\n".join(c.text for t in d.tables for row in t.rows for c in row.cells)
        return await orig(p)

    g["upload_output_file"] = cap
    outs = [await run(s, g) for s in ALL]
    return resp(outs), box["txt"]


async def test_manual_content_rules():
    st = json.loads(json.dumps(STATE))
    st["wbs"] = [dict(unit="WORKFLOW_AUTO_BUILD", quantity=1, reason="build")]
    r, txt = await _manual_text(st, n2d(fit=dict(verdict="BUILD", reasons=["ok"], kill_risks=["IT"], ask_to_see=["x"])))
    low = txt.lower()
    assert r["phase"] == "manual" and r["manual"]["url"].endswith(".docx") and r["manual"]["blueprint"][0]["step"] == "Read invoices"
    assert not any(w in low for w in ["seamless", "leverage", "ai agent", "ai assistant", "robust"])
    assert "—" not in txt
    for h in ["Recommendation", "Verdict: Build", "Who could stop this", "What to look at before building", "Proposed Auto blueprint", "The problem", "Rollout", "Indicative effort", "What we know", "Review"]:
        assert h in txt, h
    assert "STARTER RATE CARD" in txt


async def test_manual_makes_two_model_calls_and_works_from_empty_state():
    calls = []
    g, outs = await turn({"action": "manual", "user_message": "manual"}, n2d(), N4, calls)
    r = resp(outs)
    assert [c[0] for c in calls] == ["N2Out", "N4Doc"]
    assert r["manual"] and len(r["open_questions"]) > 0


async def test_manual_trigger_needs_real_words():
    _, outs = await turn({"user_message": "please write up the manual", "state_json": json.dumps(STATE)}, n2d(wants_manual=True), N4)
    assert resp(outs)["phase"] == "manual"
    _, outs = await turn({"user_message": "we use SAP too"}, n2d(wants_manual=True), N4)
    assert resp(outs)["phase"] != "manual"


async def test_attachments_degrade_instead_of_crashing():
    csv = tempfile.mktemp(suffix=".csv"); open(csv, "w").write("vendor,amount\nA,10\nB,20\n")
    g = make({"user_message": MSG, "sample_file": csv}, n2d()); await run(C.N1, g)
    assert "rows=2" in g["step_n1_profile_text"]
    empty = tempfile.mktemp(suffix=".csv"); open(empty, "w").write("")
    odd = tempfile.mktemp(suffix=".exe"); open(odd, "w").write("x")
    for ref in (empty, odd, "/nope/missing.csv"):
        g = make({"user_message": MSG, "sample_file": ref}, n2d()); await run(C.N1, g)
        assert g["step_n1_profile_text"] == "" and any("could not be used" in n for n in g["step_n1_notices"])


async def test_attachment_text_counts_as_evidence():
    csv = tempfile.mktemp(suffix=".csv"); open(csv, "w").write("vendor,amount\nA,10\n")
    _, outs = await turn({"user_message": "here is a sample", "sample_file": csv},
                         n2d(updates=[up("data_sources_and_formats", "CSV with vendor and amount", "rows=1 cols=2")]))
    assert area(resp(outs), "data_sources_and_formats")["status"] == "FACT"


async def test_weaker_never_overwrites_stronger_and_assumption_keeps_gap_open():
    _, outs = await turn({"user_message": "hmm", "state_json": json.dumps(STATE)}, n2d(updates=[up("business_goal", "something else", "", "ASSUMPTION")]))
    assert area(resp(outs), "business_goal")["value"] == "vbus"
    _, outs = await turn({"user_message": MSG}, n2d(updates=FULL[:1] + [up("current_process", "Probably manual keying", "", "ASSUMPTION")]))
    assert "confirm this assumption" in resp(outs)["reply"]


async def test_vague_answers_are_challenged():
    vague = up("volume_per_day", "a lot", "a lot of invoices", specific=False)
    st = json.loads(json.dumps(STATE)); del st["fields"]["volume_per_day"]
    _, outs = await turn({"user_message": "a lot of invoices", "state_json": json.dumps(st)}, n2d(updates=[vague]))
    r = resp(outs)
    assert area(r, "volume_per_day")["status"] == "ASSUMPTION" and "too vague" in r["reply"]


async def test_why_asking_included():
    _, outs = await turn({"user_message": "hi"}, n2d(why_asking="The steps depend on the exact flow"))
    assert "(The steps depend on the exact flow.)" in resp(outs)["reply"]


async def test_rate_card_override_and_bad_rate_card():
    n2 = n2d(wbs_units=[dict(unit="WORKFLOW_AUTO_BUILD", quantity=1, reason="b")], updates=FULL)
    full = json.dumps({k: 10 for k in RATE_KEYS})
    _, outs = await turn({"user_message": MSG, "rate_card_json": full}, n2)
    r = resp(outs)
    assert r["estimate"]["low_hours"] == 10 and r["estimate"]["calibrated"] is True and r["estimate"]["basis"] == "YOUR RATE CARD"
    _, outs = await turn({"user_message": MSG, "rate_card_json": "{oops"}, n2)
    r = resp(outs)
    assert r["estimate"]["low_hours"] == 48 and any("rate card" in n for n in r["notices"])


async def test_bad_inputs_to_helpers_do_not_crash():
    g = make({"user_message": MSG, "recent_turns": "garbage"}, n2d()); await run(C.N1, g)
    assert g["step_n1_turns"] == []


async def test_prompt_stays_small():
    calls = []
    g = make({"user_message": MSG, "state_json": json.dumps(STATE), "recent_turns": json.dumps([{"role": "user", "text": "x" * 600}] * 6)}, n2d(), calls=calls)
    for s in ALL[:2]:
        await run(s, g)
    assert calls[0][1] < 10000


import asyncio
import pytest as _pytest


def _localize():
    ns = {}
    src = C.N1.split("async def n1_localize")[0].split("step_n1_ref")[0]
    import httpx, os, tempfile
    from urllib.parse import urlparse
    ns.update(httpx=httpx, os=os, tempfile=tempfile, urlparse=urlparse)
    body = "async def n1_localize" + C.N1.split("async def n1_localize")[1].split("def n1_frame")[0]
    exec(body, ns)
    return ns["n1_localize"]


@_pytest.mark.parametrize("url", ["http://example.com/a.csv", "https://127.0.0.1/a.csv", "https://localhost/a.csv",
                                  "https://169.254.169.254/x", "file:///etc/passwd", "https://svc.internal/a.csv"])
def test_sample_fetch_refuses_unsafe_links(url):
    with _pytest.raises(ValueError):
        asyncio.run(_localize()(url))
